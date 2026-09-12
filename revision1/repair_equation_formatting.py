"""Repair math typography in the author-edited upload without rebuilding prose."""
from pathlib import Path
from copy import deepcopy
import re
import shutil
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
source = ROOT / 'revision1/FINAL/02_PeerJ_upload/cs-134580-v0.4.docx'
backup = ROOT / 'revision1/work/author_grammar_manuscript.docx'
if not backup.exists():
    shutil.copy2(source, backup)
d = Document(source)
before = [p.text for p in d.paragraphs]

def node(tag, *children):
    e = OxmlElement('m:' + tag)
    for child in children:
        e.append(child)
    return e

def run(text, upright=False):
    r = node('r')
    if upright:
        prop = node('rPr'); style = node('sty'); style.set(qn('m:val'), 'p')
        prop.append(style); r.append(prop)
    wp = OxmlElement('w:rPr'); font = OxmlElement('w:rFonts')
    font.set(qn('w:ascii'), 'Cambria Math'); font.set(qn('w:hAnsi'), 'Cambria Math')
    wp.append(font); r.append(wp)
    t = node('t'); t.text = text; t.set(qn('xml:space'), 'preserve'); r.append(t)
    return r

def sub(base, label, upright=False):
    return node('sSub', node('e', run(base, upright)), node('sub', run(label, True)))

def frac(numerator, denominator):
    return node('f', node('num', *numerator), node('den', *denominator))

def phat():
    prop = node('accPr'); char = node('chr'); char.set(qn('m:val'), '\u0302'); prop.append(char)
    return [node('acc', prop, node('e', run('p'))), run('(d)')]

equations = {
    4: [run('H(d) = −'), *phat(), sub('log','2',True), *phat(), run(' − [1 − '),
        *phat(), run('] '), sub('log','2',True), run('[1 − '), *phat(), run(']')],
    5: [run('N(d) = '), run('exp',True), run('['), frac([run('−I(d)')],[run('τ')]),
        run('],    R(d) = [1 − N(d)] H(d)')],
    7: [sub('s','norm'), run('(d) = 1 − '),
        frac([sub('rank','pool',True),run('(d) − 1')],[run('M − 1')])],
    8: [sub('s','EDE'), run('(d) = '), sub('s','norm'), run('(d) + λE(d)')]
}
repaired = []
for p in d.paragraphs:
    maths = p._p.xpath('.//m:oMath')
    match = re.search(r'\((\d)\)\s*$', p.text)
    if maths and match and int(match[1]) in equations:
        number = int(match[1]); old = maths[0]
        old.getparent().replace(old, node('oMath', *deepcopy(equations[number])))
        repaired.append(number)
    if '—' in p.text:
        assert p.text.count('—') == 2
        replacements = iter([' (', '), ' if 'limited—offline' in p.text else ') '])
        for t in p._p.xpath('.//w:t'):
            if t.text:
                t.text = ''.join(next(replacements) if c == '—' else c for c in t.text)

assert repaired == [4, 5, 7, 8], repaired
after = [p.text for p in d.paragraphs]
for a, b in zip(before, after):
    expected = a.replace('—', ' (', 1).replace('—', '), ' if 'limited—offline' in a else ') ', 1)
    assert b == expected
assert not any('—' in t.text for t in d.element.xpath('.//w:t') if t.text)
assert len(d.element.xpath('.//m:oMath')) in (8, 9)
# Section 3.4's UCB expression also needs a real square root and fraction.
ucb = 'p̂+sqrt[2 log(max(2,ΣI))/(I+1)]'
for p in d.paragraphs:
    if ucb not in p.text:
        continue
    matches = [t for t in p._p.xpath('.//w:t') if t.text and ucb in t.text]
    assert len(matches) == 1
    t = matches[0]; original_run = t.getparent()
    left, right = t.text.split(ucb)
    t.text = left
    tail = deepcopy(original_run); tail.find(qn('w:t')).text = right
    props = node('radPr'); hide = node('degHide'); hide.set(qn('m:val'), '1'); props.append(hide)
    radical = node('rad', props, node('deg'), node('e', frac(
        [run('2 '), run('log', True), run('('), run('max', True), run('(2, ΣI))')],
        [run('I + 1')])))
    inline = node('oMath', phat()[0], run(' + '), radical)
    original_run.addnext(inline); inline.addnext(tail)
assert len(d.element.xpath('.//m:oMath')) == 9
d.save(source)
shutil.copy2(source, ROOT / 'submitted_peerj_docs/cs-134580-v0.4.docx')
print('Preserved author prose; formatted equations 4, 5, 7, 8 and inline UCB; removed six em dashes.')
