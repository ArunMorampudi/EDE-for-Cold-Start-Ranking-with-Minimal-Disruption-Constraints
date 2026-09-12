"""Final numerical, provenance and document-structure audit."""
import json,re,hashlib
from pathlib import Path
from zipfile import ZipFile
import numpy as np
import pandas as pd
from docx import Document
from PIL import Image
from lxml import etree
root=Path.cwd(); result=root/'revision1/results'
test=pd.read_csv(result/'test_per_seed.csv'); val=pd.read_csv(result/'validation_per_seed.csv')
summary=pd.read_csv(result/'test_summary.csv'); replay=pd.read_csv(result/'replay/summary.csv')
assert len(test)==3900 and len(val)==100
assert set(test.seed)==set(range(1000,1020)) and set(val.seed)==set(range(100,110))
assert set(test.seed).isdisjoint(val.seed)
assert (summary.n==20).all() and len(summary)==195
assert (test.prefix_violations==0).all()
for metric in ['coverage','ndcg','ctr','query_relevant_coverage','loss_fraction','query_loss_fraction']:
    assert test[metric].between(0,1).all(),metric
assert (test.query_relevant_coverage<=test.coverage+1e-12).all()
assert test.ttf.between(0,1000).all() and test.unique.between(0,200).all()
assert set(test[test.suite=='sensitivity'].eta)=={.35,.5,.65,.8,1.}
assert set(test[test.suite=='stress'].scenario)=={'cascade','noisy','strong-bias','heterogeneous','navigational','sparse','delayed','natural'}
keys=['suite','scenario','penalty','strategy','lam','eta','m','alpha']
for _,row in summary.iterrows():
    group=test
    for key in keys: group=group[group[key]==row[key]]
    assert len(group)==20
    for metric in ['coverage','ndcg','ttf','unique','ctr']:
        assert np.isclose(group[metric].mean(),row[metric+'_mean'])
        assert np.isclose(group[metric].std(ddof=1),row[metric+'_sd'])
    if row.strategy!='Baseline':
        base=test[(test.suite==row.suite)&(test.scenario==row.scenario)&(test.penalty==row.penalty)&(test.strategy=='Baseline')].set_index('seed')
        delta=group.set_index('seed').coverage-base.coverage
        assert np.isclose(delta.mean(),row.delta_coverage)
        assert row.holm_coverage>=row.p_coverage-1e-12
allrow=replay[replay.criterion=='all'].iloc[0]
assert allrow.pages==372282
for crit in ['set_match','click_preserved']:
    assert replay[(replay.criterion==crit)&(replay.accepted==True)].pages.iloc[0]==372282
    assert replay[(replay.criterion==crit)&(replay.accepted==False)].pages.iloc[0]==0
assert allrow.delta_ctr==0 and allrow.delta_click_ndcg_ci_low<allrow.delta_click_ndcg_ci_high
source=json.loads((root/'revision1/data/source.json').read_text())
raw=root/'revision1/data/train_first_200000_sessions.txt'
assert hashlib.sha256(raw.read_bytes()).hexdigest()==source['sha256']
docs=list((root/'submitted_peerj_docs').glob('*.docx')); assert len(docs)==11
for path in docs:
    with ZipFile(path) as z:
        assert z.testzip() is None
        for name in z.namelist():
            if name.endswith('.xml'): etree.fromstring(z.read(name))
    d=Document(path)
    if re.search(r'Table[1-6]\.docx$',path.name):
        assert len(d.tables)==1, (path.name,'must contain exactly one table')
        assert not any(p.text.strip() for p in d.paragraphs)
        assert not any(p.text.strip() for s in d.sections for part in (s.header,s.footer) for p in part.paragraphs)
        assert all(s.left_margin.cm>=2.499 and s.right_margin.cm>=2.499 for s in d.sections)
    if 'tracked' not in path.name:
        assert not list(d.element.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}ins'))
        assert not list(d.element.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}del'))
manuscript=Document(root/'submitted_peerj_docs/cs-134580-v0.4.docx')
text='\n'.join(p.text for p in manuscript.paragraphs)
assert 'Figure Legends' not in text and 'Figure 1.' not in text
assert '—' not in text
original=Document(root/'revision1/original_submitted/cs-134580-v0.4.docx')
assert next(p.text for p in manuscript.paragraphs if p.text.startswith('During the preparation'))==next(p.text for p in original.paragraphs if p.text.startswith('During the preparation'))
assert len([p for p in Document(root/'submitted_peerj_docs/cs-134580-Figure-legends.docx').paragraphs if re.match(r'^Figure [1-4]\.',p.text)])==4
for n in range(1,7): assert 'Table '+str(n) in text
for n in range(1,5): assert 'Figure '+str(n) in text
maths=list(manuscript.element.iter('{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath'))
assert len(maths)==9  # Eight numbered equations and the inline UCB expression.
assert sum(bool(p._p.xpath('.//m:oMath')) and bool(re.search(r'\([1-8]\)\s*$',p.text)) for p in manuscript.paragraphs)==8
table_positions=[text.index('Table '+str(n)) for n in range(1,7)]
assert table_positions==sorted(table_positions),'Tables must first appear in numerical order'
render=json.loads((root/'revision1/qa/word/render_manifest.json').read_text(encoding='utf-8-sig'))
assert all(r['pages']==1 for r in render if re.search(r'Table[1-6]\.docx$',r['file']))
for section in manuscript.sections:
    line=section._sectPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}lnNumType')
    assert line is not None
    assert line.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}restart')=='continuous'
assert len(json.loads((root/'revision1/reviewer_response_matrix.json').read_text(encoding='utf-8')))==65
for path in (root/'submitted_peerj_docs').glob('*.png'):
    im=Image.open(path); assert all(abs(dpi-300)<1 for dpi in im.info['dpi'])
redline=json.loads((root/'revision1/qa/redline_validation.json').read_text(encoding='utf-8-sig'))
assert redline['accept_matches_clean'] and redline['reject_matches_original'] and redline['revisions']>0
report=dict(status='passed',validation_runs=100,test_runs=3900,test_conditions=195,
    replay_pages=372282,replay_sessions=200000,reviewer_comments=65,docx_files=11,numbered_equations=8,
    prefix_violations=0,tracked_changes=redline['revisions'],accept_matches_clean=True,reject_matches_original=True,
    pending_author_items=['Funding declaration','Competing-interest declaration','Original AI prompt-log reconciliation','New Zenodo version DOI'],
    limitations=['Heuristic comparators are not exact reproductions of published safe-ranking algorithms',
    'No causal or real cold-start claims from replay','Title and conclusions require author scientific review'])
(root/'revision1/validation_report.json').write_text(json.dumps(report,indent=2))
print(json.dumps(report,indent=2))
